"""Guest-list export: an Excel workbook of one event's guests.

The organizer downloads it, or has it e-mailed (to themself or to a
caterer, a venue manager…). The workbook is built server-side from the same
data the guest list shows, so both routes deliver the same file:

* « Invités » — one row per guest (deleted invitations excluded), sorted by
  table then name: name and civility, table and its note, presence, and
  invitation status and dates.
* « Tables » — every table (registry tables included, even empty) with its
  guest count, confirmed presences and guest names; « Sans table » last.
* « Résumé » — the event and its totals.

Both endpoints are organizer-scoped (another organizer's event is a 404).
The e-mail route reuses the invitation mailer: same provider, same
« never report a send that did not happen » contract, same per-account
throttle.
"""
from __future__ import annotations

import io
import logging
import re
import unicodedata
from datetime import datetime
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from django.conf import settings
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils.timezone import now as timezone_now
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from rest_framework import serializers, status
from rest_framework.exceptions import NotFound
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.views import APIView

from events.models import EventModel
from invitations.mailer import Attachment, DeliveryError, send_invitation_email
from invitations.models import Invitation
from invitations.sharing import (
    EmailDeliveryFailed,
    EmailDeliveryUnavailable,
    french_date,
    french_time,
)

logger = logging.getLogger(__name__)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

INK = "2B2620"
GOLD = "B08A45"
HEADER_FILL = PatternFill("solid", fgColor="2B2620")
BAND_FILL = PatternFill("solid", fgColor="F7F1E6")
CONFIRMED_FILL = PatternFill("solid", fgColor="E8F3EC")
THIN = Side(style="thin", color="E3D6BE")

STATUS_LABELS = {
    Invitation.State.ACTIVE: "Active",
    Invitation.State.REVOKED: "Révoquée",
    Invitation.State.EXPIRED: "Expirée",
}
NO_TABLE = "Sans table"


# --------------------------------------------------------------- helpers


def _event_zone(event: EventModel):
    try:
        return ZoneInfo(event.timezone or "Africa/Kinshasa")
    except (ZoneInfoNotFoundError, ValueError):
        return ZoneInfo("Africa/Kinshasa")


def _local(value: datetime | None, zone) -> datetime | None:
    """Aware datetime → naive local time (Excel stores no time zone)."""
    if value is None:
        return None
    return value.astimezone(zone).replace(tzinfo=None, microsecond=0)


def _table_sort_key(number: str):
    """Natural order: 2 before 10, labelled tables after numbers, none last."""
    if not number:
        return (2, 0, "")
    match = re.match(r"^\s*(\d+)", number)
    if match:
        return (0, int(match.group(1)), number.casefold())
    return (1, 0, number.casefold())


def _put(sheet: Worksheet, row: int, column: int, value, **style):
    """Write a cell; text is always stored as text.

    openpyxl turns any string starting with « = » into a live formula, and a
    guest name or a table note is user input: forcing the string type blocks
    spreadsheet formula injection.
    """
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    for key, item in style.items():
        setattr(cell, key, item)
    return cell


def _header_row(sheet: Worksheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        _put(
            sheet,
            row,
            column,
            label,
            font=Font(bold=True, color="FFFFFF"),
            fill=HEADER_FILL,
            alignment=Alignment(vertical="center", wrap_text=True),
            border=Border(bottom=THIN),
        )
    sheet.row_dimensions[row].height = 30


def _fit_columns(sheet: Worksheet, widths: dict[int, float], *, first_row: int) -> None:
    """Column widths from the longest value, within sensible bounds."""
    for column in range(1, sheet.max_column + 1):
        if column in widths:
            sheet.column_dimensions[get_column_letter(column)].width = widths[column]
            continue
        longest = 0
        for (value,) in sheet.iter_rows(
            min_row=first_row, min_col=column, max_col=column, values_only=True
        ):
            if value is None:
                continue
            text = value.strftime("%d/%m/%Y %H:%M") if isinstance(value, datetime) else str(value)
            longest = max(longest, max((len(line) for line in text.splitlines()), default=0))
        sheet.column_dimensions[get_column_letter(column)].width = max(10, min(48, longest + 2))


def _title_block(sheet: Worksheet, event: EventModel, title: str, exported_at: datetime) -> int:
    """Event heading above a sheet's table; returns the next free row."""
    _put(sheet, 1, 1, title, font=Font(bold=True, size=14, color=INK))
    when = " à ".join(filter(None, [french_date(event.event_date), french_time(event.event_time)]))
    place = " — ".join(filter(None, [event.venue_name, event.venue_address]))
    _put(sheet, 2, 1, event.title, font=Font(bold=True, size=12, color=GOLD))
    _put(sheet, 3, 1, " · ".join(filter(None, [when, place])), font=Font(italic=True, color="7A6A50"))
    _put(
        sheet,
        4,
        1,
        f"Exporté le {exported_at.strftime('%d/%m/%Y à %H:%M')} depuis MalaikaEvents",
        font=Font(size=9, color="9A8B73"),
    )
    return 6


# -------------------------------------------------------------- workbook


def guest_rows(event: EventModel):
    """The event's live invitations, in seating order."""
    invitations = list(
        Invitation.objects.filter(event_model=event)
        .exclude(state=Invitation.State.DELETED)
        .select_related("table")
    )
    invitations.sort(key=lambda item: (_table_sort_key(item.table_number), item.guest_name.casefold()))
    return invitations


def build_guest_workbook(event: EventModel) -> bytes:
    zone = _event_zone(event)
    exported_at = _local(timezone_now(), zone)
    invitations = guest_rows(event)

    workbook = Workbook()

    # ---- Invités
    sheet = workbook.active
    sheet.title = "Invités"
    row = _title_block(sheet, event, "Liste des invités", exported_at)
    labels = [
        "N°",
        "Invité",
        "Civilité",
        "Nom",
        "Table",
        "Précision de table",
        "Présence",
        "Confirmée le",
        "Invitation",
        "Émise le",
        "Expire le",
    ]
    header_row = row
    _header_row(sheet, header_row, labels)
    civilities = dict(Invitation.Civility.choices)
    for index, invitation in enumerate(invitations, start=1):
        row += 1
        confirmed = invitation.presence_confirmed_at is not None
        values = [
            index,
            invitation.display_name,
            "" if invitation.civility == Invitation.Civility.NONE else civilities.get(invitation.civility, ""),
            invitation.guest_name,
            invitation.table_number or "",
            invitation.table_text or "",
            "Confirmée" if confirmed else "En attente",
            _local(invitation.presence_confirmed_at, zone),
            STATUS_LABELS.get(invitation.status, invitation.status),
            _local(invitation.issued_at, zone),
            _local(invitation.expires_at, zone),
        ]
        for column, value in enumerate(values, start=1):
            cell = _put(sheet, row, column, value, border=Border(bottom=THIN))
            if isinstance(value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
        presence = sheet.cell(row=row, column=7)
        if confirmed:
            presence.fill = CONFIRMED_FILL
            presence.font = Font(bold=True, color="1F6B3A")
    if invitations:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(labels))}{row}"
    else:
        _put(sheet, header_row + 1, 1, "Aucun invité pour le moment.", font=Font(italic=True, color="7A6A50"))
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=3)
    _fit_columns(sheet, {1: 6}, first_row=header_row)
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    # ---- Tables
    tables = workbook.create_sheet("Tables")
    row = _title_block(tables, event, "Plan de table", exported_at)
    table_header = row
    _header_row(tables, table_header, ["Table", "Précision", "Invités", "Présences confirmées", "Noms des invités"])
    seating: dict[str, dict] = {}
    for table in event.tables.all():
        seating.setdefault(table.number, {"text": table.text, "guests": [], "confirmed": 0})
    for invitation in invitations:
        if invitation.status == Invitation.State.REVOKED:
            continue  # a revoked guest no longer holds a seat
        key = invitation.table_number or NO_TABLE
        entry = seating.setdefault(key, {"text": invitation.table_text, "guests": [], "confirmed": 0})
        if not entry["text"] and invitation.table_text:
            entry["text"] = invitation.table_text
        entry["guests"].append(invitation.display_name)
        entry["confirmed"] += invitation.presence_confirmed_at is not None
    ordered = sorted(seating, key=lambda key: _table_sort_key("" if key == NO_TABLE else key))
    for number in ordered:
        entry = seating[number]
        row += 1
        values = [number, entry["text"] or "", len(entry["guests"]), entry["confirmed"], "\n".join(entry["guests"])]
        for column, value in enumerate(values, start=1):
            _put(
                tables,
                row,
                column,
                value,
                border=Border(bottom=THIN),
                alignment=Alignment(vertical="top", wrap_text=column == 5),
            )
        if number == NO_TABLE:
            tables.cell(row=row, column=1).font = Font(italic=True, color="7A6A50")
    if not ordered:
        _put(tables, table_header + 1, 1, "Aucune table pour le moment.", font=Font(italic=True, color="7A6A50"))
    tables.freeze_panes = tables.cell(row=table_header + 1, column=1)
    _fit_columns(tables, {5: 48}, first_row=table_header)

    # ---- Résumé
    summary = workbook.create_sheet("Résumé")
    row = _title_block(summary, event, "Résumé", exported_at)
    statuses = [invitation.status for invitation in invitations]
    seated = [i for i in invitations if i.table_number and i.status != Invitation.State.REVOKED]
    figures = [
        ("Invités", len(invitations)),
        ("Présences confirmées", sum(1 for i in invitations if i.presence_confirmed_at)),
        ("En attente de confirmation", sum(1 for i in invitations if not i.presence_confirmed_at)),
        ("Invitations actives", statuses.count(Invitation.State.ACTIVE)),
        ("Invitations expirées", statuses.count(Invitation.State.EXPIRED)),
        ("Invitations révoquées", statuses.count(Invitation.State.REVOKED)),
        ("Tables", len([key for key in seating if key != NO_TABLE])),
        ("Invités placés à une table", len(seated)),
    ]
    _header_row(summary, row, ["Indicateur", "Valeur"])
    for label, value in figures:
        row += 1
        _put(summary, row, 1, label, border=Border(bottom=THIN))
        _put(summary, row, 2, value, border=Border(bottom=THIN), font=Font(bold=True, color=INK))
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 14

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False

    workbook.properties.title = f"Invités — {event.title}"
    workbook.properties.creator = "MalaikaEvents"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(event: EventModel) -> str:
    """« invites-mariage-de-grace-et-eric-2026-12-12.xlsx » (ASCII, safe)."""
    ascii_title = unicodedata.normalize("NFKD", event.title).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_title.lower()).strip("-")[:60] or "evenement"
    return f"invites-{slug}-{event.event_date.isoformat()}.xlsx"


# ----------------------------------------------------------------- views


def _owned_event(request, pk) -> EventModel:
    event = EventModel.objects.filter(pk=pk, organizer=request.user).first()
    if event is None:
        raise NotFound("Ressource introuvable.")
    return event


class GuestListExportView(APIView):
    """``GET /api/events/{pk}/guests/export/`` — the workbook as a download."""

    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        event = _owned_event(request, pk)
        content = build_guest_workbook(event)
        response = HttpResponse(content, content_type=XLSX_CONTENT_TYPE)
        response["Content-Disposition"] = f'attachment; filename="{export_filename(event)}"'
        response["Cache-Control"] = "no-store"
        return response


class GuestListEmailSerializer(serializers.Serializer):
    recipient = serializers.EmailField(max_length=254)
    message = serializers.CharField(required=False, allow_blank=True, max_length=1500, default="")


class GuestListEmailView(APIView):
    """``POST /api/events/{pk}/guests/export/email/`` — e-mail the workbook."""

    permission_classes = [IsAuthenticated]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "share-email"

    def post(self, request, pk):
        event = _owned_event(request, pk)
        serializer = GuestListEmailSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        if not getattr(settings, "INVITATION_EMAIL_ENABLED", False):
            raise EmailDeliveryUnavailable()

        recipient = serializer.validated_data["recipient"]
        note = serializer.validated_data.get("message", "").strip()
        count = Invitation.objects.filter(event_model=event).exclude(state=Invitation.State.DELETED).count()
        site_url = getattr(settings, "SITE_URL", "https://www.malaikaevent.com")
        organizer_name = (getattr(request.user, "full_name", "") or "").strip() or request.user.email
        context = {
            "subject": f"Liste des invités — {event.title}",
            "event_title": event.title,
            "event_when": " à ".join(filter(None, [french_date(event.event_date), french_time(event.event_time)])),
            "venue_name": event.venue_name,
            "guest_count": count,
            "note_paragraphs": [block.strip() for block in note.split("\n\n") if block.strip()],
            "note": note,
            "organizer": organizer_name,
            "filename": export_filename(event),
            "site_url": site_url,
            "site_label": site_url.split("://")[-1],
        }
        try:
            send_invitation_email(
                to=recipient,
                subject=context["subject"],
                html=render_to_string("invitations/email/guest_list.html", context),
                text=render_to_string("invitations/email/guest_list.txt", context),
                reply_to=request.user.email or None,
                attachments=[
                    Attachment(export_filename(event), build_guest_workbook(event), XLSX_CONTENT_TYPE)
                ],
            )
        except DeliveryError:  # never report a send that did not happen
            logger.warning("Guest list e-mail failed (event=%s)", event.pk)
            raise EmailDeliveryFailed()

        logger.info("Guest list e-mail accepted by provider (event=%s)", event.pk)
        return Response({"status": "accepted"}, status=status.HTTP_200_OK)
